#!/usr/bin/env python3
"""Bench results → one Excel workbook to explore by hand (2026-10-21).

  python3 infra/bench/bench_xlsx.py BENCH_RESULTS_DIR [-o out.xlsx]

BENCH_RESULTS_DIR is what ops/bench.yml fetches (<repo>/bench-results/): one folder per plan run (p1-maker-smoke, …)
with throughput.csv, state.json, results.jsonl, report.md, up-*.log, and boxstats[-smoke].csv beside them.
Variant attributes (engine, slots, context, KV, streaming, MTP, GPUs) come from the plan files in infra/bench/plans/.

Sheets (2026-10-21, user: readable names): Read me (sheets, terms, column names) · Summary (one row per configuration,
formulas, charts) · Load levels (the main table) · Configurations · Queue items · Answer tasks · Box per item · Raw load
levels / requests / box samples / events (the files unchanged, readable headers). Every data sheet is an Excel table.
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import re
import statistics
import sys
import tomllib
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.data_source import NumDataSource, NumRef
from openpyxl.chart.error_bar import ErrorBars
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).resolve().parent
FONT = "Arial"
HEAD_FILL = PatternFill("solid", start_color="1F3864")


def num(x):
    if x in (None, "", "None", "–"):
        return None
    try:
        f = float(x)
        return int(f) if f.is_integer() and "." not in str(x) else f
    except (TypeError, ValueError):
        return x


# ------------------------------------------------------------------------------------------- variant attributes
def plan_variants() -> dict:
    out = {}
    for f in sorted((HERE / "plans").glob("*.toml")):
        try:
            p = tomllib.loads(f.read_text())
        except tomllib.TOMLDecodeError:
            continue
        for v in p.get("variant", []):
            out[v["id"]] = (p["name"], {"gpus": p.get("gpus"), **v})       # plan-level GPU set unless the variant has one
    return out


def attrs(vid: str, v: dict | None) -> dict:
    v = v or {}
    sa = {**v.get("args_set", {}), **v.get("strata_args_set", {})}
    cmd = v.get("command") or []
    svc = v.get("service", "")
    if "llama.cpp" in v.get("image", ""):         # llama.cpp in another engine's slot (the sub-maker B grid)
        engine, model = "llama.cpp", "27B GSQ-RCO IQ3_S"
    elif svc == "strata":
        engine, model = "Strata", "Flash-Next IQ3_S"
    elif svc == "vllm-sub":
        engine, model = "vLLM", "27B AWQ-INT4"
    elif svc == "llama-guard":
        engine, model = "llama.cpp", "Ministral 3B (guard)"
    else:
        engine, model = svc, ""
    role = ("maker (scenario B)" if vid.startswith("makerB") else "maker" if vid.startswith(("maker", "ref-maker"))
            else "guard" if vid.startswith("guard") else "sub-maker")
    kv = sa.get("--kv") or sa.get("--kv-cache-dtype") or (cmd[cmd.index("-ctk") + 1] if "-ctk" in cmd else None)
    if not kv:                                     # the production setting the variant inherits
        kv = {"strata": "int8", "vllm-sub": "fp8", "llama-guard": "q8_0"}.get(svc)
    mtp = ("draft-mtp" in cmd if cmd else
           ("--mtp" not in v.get("strata_args_remove", [])) if svc == "strata"
           else any("speculative" in a for a in v.get("args_add", [])))
    # cards the engine runs on: the variant's own device list (compose override), else the slot's production default
    devs = [d for r in (((v.get("compose") or {}).get("deploy") or {}).get("resources") or {}).get("reservations", {})
            .get("devices", []) for d in r.get("device_ids", [])]
    gpus = devs or {"strata": [0, 1], "vllm-sub": [2], "llama-guard": [1]}.get(svc) or v.get("gpus")
    a = {"role": role, "engine": engine, "model": model, "slot": svc, "slots": v.get("slots"),
         "ctx_k": (v.get("ctx") or 0) // 1024 or None, "kv": kv,
         "kv_streaming": "--kv-resident" in sa, "mtp": bool(mtp), "gpus": len(gpus) if gpus else None,
         "vision_off": any("limit-mm-per-prompt" in x for x in v.get("args_add", [])),
         "window_k": int(sa["--kv-resident"]) // 1024 if "--kv-resident" in sa else None,
         "fallback": v.get("only_if_failed")}
    a.update(describe(a, gpus))
    if vid.startswith("ref-"):                     # the exclusive block's isolated references: nothing else running
        a["configuration"] += " — alone on the box (reference)"
        a["short_name"] += " alone"
    return a


SLOT_CARD = {"strata": "#1+#2", "llama-guard": "#2", "vllm-sub": "#3"}
OWN_SLOT = {"Strata": "strata", "vLLM": "vllm-sub"}


def describe(a, gpus):
    """A configuration in words (long) and as a chart label (short)."""
    cards = "+".join(f"#{g + 1}" for g in sorted(int(x) for x in gpus)) if gpus else "?"
    kv = a["kv"] or "?"
    stream = f", KV streamed to RAM (window {a['window_k'] or 32}K)" if a["kv_streaming"] else ""
    foreign = a["engine"] == "llama.cpp" and a["role"] == "sub-maker" and a["slot"] != "vllm-sub"
    long = (f"{a['role'].capitalize()}: {a['model']} on {a['engine']}, {a['slots']} parallel × {a['ctx_k']}K context, "
            f"KV cache {kv}{stream}, MTP {'on' if a['mtp'] else 'off'}{', vision off' if a['vision_off'] else ''}, "
            f"card {cards}" + (f" (in the {a['slot']} slot)" if foreign else "")
            + (f" — only if {a['fallback']} did not start" if a["fallback"] else ""))
    r = {"maker": "Maker", "maker (scenario B)": "Maker B", "sub-maker": "Sub", "guard": "Guard"}.get(a["role"], a["role"])
    short = (f"{r} {a['engine']} {a['slots']}×{a['ctx_k']}K {kv}"
             + (f" stream{a['window_k'] or 32}K" if a["kv_streaming"] else "") + (" MTP" if a["mtp"] else "")
             + (" novision" if a["vision_off"] else "") + f" {cards}")
    return {"configuration": long, "short_name": short}


def test_name(workload, fill):
    if workload == "depth":
        return f"Long prompt, {fill * 100:.0f} % of the context" if fill is not None else "Long prompt"
    return {"ramp": "Rising number of parallel requests", "answerlen": "Answer-length tasks",
            "guardload": "Guard load", "agents": "Agent loops", "sustained": "Sustained load",
            "effort": "Effort tasks"}.get(workload, workload)


def variant_of(item_or_label: str) -> str:
    return item_or_label.split("/", 1)[0]


# ------------------------------------------------------------------------------------------- readers
def read_box(results: Path) -> list[list]:
    """boxstats*.csv → long rows [ts, key, value]; tolerant of a file that is not a CSV."""
    rows = []
    for f in results.glob("boxstats*.csv"):
        raw = f.read_bytes()
        if raw[:2] == b"\x1f\x8b":            # not a CSV (e.g. an archive saved under that name)
            continue
        for r in csv.reader(io.StringIO(raw.decode("utf-8", "replace"))):
            if len(r) == 3:
                try:
                    rows.append([float(r[0]), r[1], float(r[2])])
                except ValueError:
                    pass
    return rows


def box_from_report(report: Path) -> list[dict]:
    """Fallback: the report's per-item box table."""
    text = report.read_text(errors="replace")
    m = re.search(r"## Box during each item[^\n]*\n\n(\|.*?)(\n\n|\Z)", text, re.S)
    if not m:
        return []
    lines = [l for l in m.group(1).splitlines() if l.startswith("|")]
    head = [h.strip() for h in lines[0].strip("|").split("|")]
    out = []
    for l in lines[2:]:
        cells = [c.strip() for c in l.strip("|").split("|")]
        out.append({h: num(c) for h, c in zip(head, cells)})
    return out


def start_info(log: Path) -> dict:
    t = log.read_text(errors="replace")
    g = lambda pat: (re.search(pat, t) or [None, None])[1]   # noqa: E731
    ready = g(r"ready after (\d+)s")
    fail = None if ready else next((l.split("|", 1)[-1].strip() for l in t.splitlines()
                                    if re.search(r"needs|error|failed|not enough|out of memory", l, re.I)),
                                   t.splitlines()[0] if t.strip() else "")
    return {"ready_s": num(ready), "started": bool(ready), "fail_reason": fail,
            "strata_experts_in_vram": num(g(r"expert cache \((\d+) experts")),
            "strata_expert_vram_gib": num(g(r"experts, ([\d.]+) GiB of VRAM")),
            "strata_experts_ram_gib": num(g(r"experts loaded: ([\d.]+) GiB")),
            "strata_expert_load_gibs": num(g(r"at ([\d.]+) GiB/s")),
            "vllm_kv_tokens": num((g(r"GPU KV cache size: ([\d,]+) tokens") or "").replace(",", "") or None),
            "vllm_max_concurrency": num(g(r"Maximum concurrency for [\d,]+ tokens per request: ([\d.]+)x")),
            "vllm_kv_gib": num(g(r"Available KV cache memory: ([\d.]+) GiB")),
            "llama_slots": num(g(r"n_slots = (\d+)")), "llama_ctx_slot": num(g(r"n_ctx_slot = (\d+)"))}


# ------------------------------------------------------------------------------------------- readable names
# every column as it is labelled in the workbook (raw sheets too: same values, readable headers); "Read me" lists them
LABEL = {
    "run": "Queue", "plan": "Plan", "variant": "Configuration id", "configuration": "Configuration",
    "short_name": "Short name", "workload": "Test", "test": "Test", "fill": "Prompt size (share of context)",
    "clients": "Parallel requests", "streams": "Parallel requests", "role": "Role", "engine": "Engine", "model": "Model",
    "slot": "Runs in the slot of", "slots": "Parallel slots", "ctx_k": "Context per slot (K tokens)",
    "kv": "KV cache format", "kv_streaming": "KV streamed to RAM", "window_k": "Streaming window (K tokens)",
    "mtp": "MTP drafting", "vision_off": "Vision off", "gpus": "Cards", "fallback": "Runs only if this did not start",
    "prompt_tokens": "Prompt length (tokens)", "max_tokens": "Answer length (tokens)",
    "ttft_s_min": "Wait for first token, fastest (s)", "ttft_s_med": "Wait for first token, median (s)",
    "ttft_s_max": "Wait for first token, slowest (s)",
    "prefill_tps_slot": "Prompt reading per request (tok/s)", "prefill_tps_agg": "Prompt reading, all requests (tok/s)",
    "decode_tps_steady": "Writing speed per request, steady (tok/s)",
    "decode_tps_med": "Writing speed per request, median (tok/s)", "decode_tps_min": "Writing speed, slowest request (tok/s)",
    "aggregate_tps": "Writing speed, all requests (tok/s)", "wall_s": "Duration (s)",
    "cache_hit": "Prompt from cache (share)", "spec_accept": "MTP drafts accepted (share)",
    "kv_usage_max": "KV cache fill, peak (share)", "waiting_max": "Requests waiting, peak",
    "preemptions": "Requests pushed out (vLLM)", "expert_hit": "Strata experts found on the card (share)",
    "pcie_share": "Strata experts fetched over PCIe (share)", "requests": "Requests", "errors": "Errors",
    "prompt_source": "Prompt text", "chars_per_token": "Characters per token", "first_error": "First error",
    "label": "Queue item", "t_start": "Start (epoch s)", "t_end": "End (epoch s)",
    "started": "Started", "ready_s": "Start time (s)", "fail_reason": "Why it did not start",
    "strata_experts_in_vram": "Strata experts on the cards", "strata_expert_vram_gib": "Strata expert cache on the cards (GiB)",
    "strata_experts_ram_gib": "Strata experts in RAM (GiB)", "strata_expert_load_gibs": "Strata expert loading (GiB/s)",
    "vllm_kv_tokens": "vLLM KV room (tokens)", "vllm_kv_gib": "vLLM KV room (GiB)",
    "vllm_max_concurrency": "vLLM full-length requests that fit", "llama_slots": "llama.cpp slots",
    "llama_ctx_slot": "llama.cpp context per slot (tokens)",
    "item": "Queue item", "status": "Status", "duration_s": "Duration (s)", "attempts": "Attempts",
    "finished": "Finished (epoch s)", "error": "Error", "s": "Duration (s)",
    "task": "Task", "shape": "Task shape", "level": "Effort level", "rep": "Repeat", "ok": "Request ok",
    "passed": "Answer passed the check", "finish": "Why it stopped", "completion_tokens": "Answer tokens",
    "reasoning_tokens": "Thinking tokens", "reasoning_src": "Thinking counted from", "cached_tokens": "Prompt tokens from cache",
    "ttft_s": "Wait for first token (s)", "think_s": "Thinking time (s)", "total_s": "Total time (s)", "attempt": "Attempt",
    "epoch": "Epoch (s)", "time": "Time", "event": "Event",
}


def label(k: str) -> str:
    if k in LABEL:
        return LABEL[k]
    m = re.match(r"gpu(\d+)\.(\w+)$", k)
    if m:
        what = {"util": "busy (%)", "mem_mb": "memory used (MB)", "power_w": "power (W)", "temp_c": "temperature (°C)",
                "sm_mhz": "clock (MHz)", "pcie_rx_mbs": "PCIe into the card (MB/s)",
                "pcie_tx_mbs": "PCIe out of the card (MB/s)"}.get(m.group(2), m.group(2))
        return f"Card #{int(m.group(1)) + 1} {what}"
    m = re.match(r"disk\.([\w-]+)\.(read|write)_mbs$", k)
    if m:
        return f"Disk {m.group(1)} {m.group(2)} (MB/s)"
    return {"cpu.busy_pct": "CPU busy (%)", "cpu.iowait_pct": "CPU waiting for disks (%)",
            "cpu.max_core_pct": "CPU busiest core (%)", "cpu.cores_over_90": "CPU cores above 90 %", "ram.used_gb": "RAM used (GB)",
            "ram.avail_gb": "RAM free (GB)", "ram.cached_gb": "RAM file cache (GB)", "swap.used_gb": "Swap used (GB)",
            "membw.read_mbs": "RAM bandwidth read (MB/s)", "membw.write_mbs": "RAM bandwidth write (MB/s)"}.get(k, k)


GLOSSARY = [
    ("Configuration", "One engine setup that was benchmarked: model, engine, how many requests it serves in parallel "
                      "(slots), context per slot, KV cache format, MTP, cards. The id (e.g. maker-b2-128k-int8-stream) "
                      "is the name in the plan files; \"Configuration\" says the same in words."),
    ("Maker / Sub / Guard", "The roles: the maker (Strata, Qwen3.8-Flash-Next) does the main work; sub-makers "
                            "(Qwen3.8-27B on vLLM or llama.cpp) take delegated parts; the guard checks tool results."),
    ("Prompt reading", "How fast the engine processes the prompt before the first answer token (prefill), in tokens/s."),
    ("Writing speed", "How fast answer tokens come out (decode), in tokens/s. \"Steady\" counts only the time when every "
                      "request of a round is writing — no other request's prompt reading in between; the median "
                      "includes that overlap and is lower when requests queue."),
    ("Wait for first token", "Time from sending a request to its first answer token (TTFT): prompt reading plus any "
                             "time waiting in the queue. Slowest minus fastest ≈ the queue."),
    ("Parallel requests / slots", "Requests sent at the same time vs. how many the engine serves at once. More requests "
                                  "than slots = some wait (the \"+2\" levels)."),
    ("Long prompt, N % of the context", "Prompts filling that share of one slot's context (e.g. 90 % of 128K ≈ 115K "
                                        "tokens) — how the engine behaves deep into a conversation."),
    ("KV cache", "The engine's memory of the conversation so far. int8 / fp8 / q8_0 = 8 bits per value; k8v4 = 8-bit "
                 "keys, 4-bit values (smaller); f16/auto = 16 bits. \"Streamed to RAM\": Strata keeps the whole KV in "
                 "system RAM and only a window on the card, so the card keeps more room for experts."),
    ("MTP", "Multi-token prediction: a small extra layer guesses the next tokens, the model checks them in one step. "
            "\"Drafts accepted\" = share of guesses that were right; higher means faster writing."),
    ("Strata experts on the card", "Strata keeps all expert weights in RAM and the most used ones on the cards; "
                                   "\"found on the card\" = share of expert uses served from the card."),
    ("Vision off", "vLLM started without the image encoder's memory reserve — that memory goes to the KV cache."),
    ("Cards", "#1, #2, #3 = the three RTX 3090 (PCIe slots x16_1, x16_2, x16_3)."),
]


# ------------------------------------------------------------------------------------------- workbook helpers
def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 32


def data_sheet(wb, name, cols, rows, widths=None, notes=None, table=None):
    """A sheet with readable headers (label()), the values as they are, as an Excel table (Insert → PivotTable)."""
    ws = wb.create_sheet(name)
    ws.append([label(c) for c in cols])
    for r in rows:
        ws.append([r.get(c) for c in cols])
    style_header(ws, len(cols))
    ws.row_dimensions[1].height = 45
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT)
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(c, max(11, min(24, len(label(c)) // 2 + 4)))
    if rows:
        ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
        tb = Table(displayName=table or re.sub(r"\W", "", name), ref=ref)
        tb.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tb)
    if notes:
        r0 = len(rows) + 3
        for i, n in enumerate(notes):
            ws.cell(row=r0 + i, column=1, value=n).font = Font(name=FONT, italic=True, color="555555")
    return ws


ATTR_COLS = ["role", "engine", "model", "slots", "ctx_k", "kv", "kv_streaming", "window_k", "mtp", "vision_off", "gpus",
             "slot"]
METRIC_COLS = ["prompt_tokens", "max_tokens", "ttft_s_min", "ttft_s_med", "ttft_s_max", "prefill_tps_slot",
               "prefill_tps_agg", "decode_tps_steady", "decode_tps_med", "decode_tps_min", "aggregate_tps", "wall_s",
               "cache_hit", "spec_accept", "kv_usage_max", "waiting_max", "preemptions", "expert_hit", "pcie_share",
               "requests", "errors", "prompt_source", "chars_per_token", "first_error"]
LEVEL_COLS = ["short_name", "variant", "test", "fill", "clients"] + METRIC_COLS + ATTR_COLS + ["run", "configuration"]
# Summary: (heading, formula kind, Load-levels column, test filter)
SUMMARY = [
    ("Prompt reading per request (tok/s), average of all tests", "avg", "prefill_tps_slot", None),
    ("Writing speed per request, steady (tok/s), average", "avg", "decode_tps_steady", None),
    ("Writing speed, all requests (tok/s), best level", "max", "aggregate_tps", None),
    ("Long prompt 90 %: prompt reading per request (tok/s)", "avg", "prefill_tps_slot", "Long prompt, 90 % of the context"),
    ("Long prompt 90 %: wait for first token, slowest (s)", "avg", "ttft_s_max", "Long prompt, 90 % of the context"),
    ("Long prompt 90 %: writing speed per request, steady (tok/s)", "avg", "decode_tps_steady",
     "Long prompt, 90 % of the context"),
    ("MTP drafts accepted (share), average", "avg", "spec_accept", None),
    ("KV cache fill, peak (share)", "max", "kv_usage_max", None),
    ("Requests pushed out (vLLM), total", "sum", "preemptions", None),
    ("Strata experts found on the card (share)", "avg", "expert_hit", None),
    ("Errors, total", "sum", "errors", None),
]
SHEETS = [
    ("Summary", "One row per configuration: what it is, how fast it started, and the key numbers (formulas over "
                "\"Load levels\"), with charts below."),
    ("Load levels", "The main table: one row per measured load level — configuration, test, parallel requests and "
                    "every number measured. Use this for your own pivots."),
    ("Configurations", "Every configuration with its settings and what the engine reported at start (or why it did "
                       "not start)."),
    ("Queue items", "Every queue item (configuration × test): status, duration, error."),
    ("Answer tasks", "The answer-length tasks: one row per request (tokens, thinking, time, passed)."),
    ("Box per item", "The box during each queue item: cards, CPU, RAM, disks averaged over the item's time (memory = "
                     "peak). Queues ran in parallel, so an item's numbers include the other queues' load."),
    ("Raw load levels", "throughput.csv of every queue, unchanged (readable headers)."),
    ("Raw requests", "results.jsonl of every queue, unchanged (readable headers)."),
    ("Raw box samples", "boxstats.csv, one row per 5-second sample, one column per measurement."),
    ("Raw events", "events.log of every queue: starts, stops, skips, failures."),
    ("Charts", "Only charts: per role, prompt reading and writing speed — bar = average, whisker = slowest … fastest "
               "level; colour family = slots × context."),
    ("Chart data", "The numbers behind the charts (formulas over \"Load levels\")."),
]


# ------------------------------------------------------------------------------------------- Charts (user 2026-10-21)
# one chart per role and measure; bars = average over the configuration's load levels, whiskers = min … max; colour
# family = the shape (parallel slots × context), shades = the configurations of that shape, side by side
FAMILIES = {"green": ["1B5E20", "2E7D32", "43A047", "66BB6A", "A5D6A7", "C8E6C9"],
            "blue": ["0D47A1", "1565C0", "1E88E5", "42A5F5", "90CAF9", "BBDEFB"],
            "orange": ["E65100", "EF6C00", "FB8C00", "FFA726", "FFCC80", "FFE0B2"],
            "purple": ["4A148C", "6A1B9A", "8E24AA", "AB47BC", "CE93D8", "E1BEE7"],
            "teal": ["004D40", "00695C", "00897B", "26A69A", "80CBC4", "B2DFDB"],
            "red": ["B71C1C", "C62828", "E53935", "EF5350", "EF9A9A", "FFCDD2"],
            "brown": ["3E2723", "4E342E", "6D4C41", "8D6E63", "BCAAA4", "D7CCC8"],
            "pink": ["880E4F", "AD1457", "D81B60", "EC407A", "F48FB1", "F8BBD0"],
            "grey": ["212121", "424242", "616161", "9E9E9E", "BDBDBD", "E0E0E0"]}
PREFERRED = {"2×128K": "green", "4×32K": "blue", "1×128K": "teal", "2×64K": "orange", "1×64K": "purple",
             "2×32K": "red", "1×32K": "brown", "6×16K": "green", "3×32K": "blue", "4×16K": "teal", "3×16K": "orange"}
CHART_ROLES = [("Maker — Strata", lambda v: v["role"] == "maker" and v["engine"] == "Strata"),
               ("Maker — Strata vs vLLM TP2 (scenario B)", lambda v: v["role"].startswith("maker") and
                (v["engine"] != "Strata" or v["slots"] == 4 and v["ctx_k"] == 64 or v["slots"] == 2 and v["ctx_k"] == 64)),
               ("Sub-maker — vLLM AWQ and llama.cpp IQ3_S", lambda v: v["role"] == "sub-maker"),
               ("Guard — llama.cpp", lambda v: v["role"] == "guard")]
CHART_MEASURES = [("prefill_tps_slot", "prompt reading per request (tok/s)"),
                  ("decode_tps_steady", "writing speed per request, steady (tok/s)")]


def shape_of(v):
    s = f"{v['slots']}×{v['ctx_k']}K"
    return s + (f" on {v['gpus']} cards" if v["role"] == "maker" and (v["gpus"] or 0) > 2 else "") + \
        (" (vLLM TP2)" if v["role"] == "maker (scenario B)" else "")


def charts_sheet(wb, vrows, levels, rng):
    """'Chart data' holds one block per chart (formulas over 'Load levels'); 'Charts' holds only the charts."""
    data = wb.create_sheet("Chart data")
    cs = wb.create_sheet("Charts", 1)                # right after the Read me
    cs["A1"] = "Charts — bar = average over the configuration's load levels, whisker = slowest … fastest level"
    cs["A1"].font = Font(name=FONT, bold=True, size=14)
    cs["A2"] = ("Colour family = shape (parallel slots × context), shades = the configurations of that shape next to "
                "each other. Values: sheet \"Chart data\" (formulas over \"Load levels\").")
    cs["A2"].font = Font(name=FONT, italic=True, color="555555")
    row, chart_row = 1, 4
    for title, pick in CHART_ROLES:
        members = [v for v in vrows if pick(v)]
        for m, mname in CHART_MEASURES:
            have = {l["short_name"] for l in levels if isinstance(l.get(m), (int, float))}
            ms = [v for v in members if v["short_name"] in have]
            if len(ms) < 2:
                continue
            shapes = sorted({shape_of(v) for v in ms},
                            key=lambda s: (int(s.split("×")[0]) * int(s.split("×")[1].split("K")[0]), s))
            fam, used = {}, set()
            for s in shapes:                          # the preferred family if still free, else the next free one
                f = PREFERRED.get(s.split(" ")[0]) if s.split(" ")[0] == s else None
                if not f or f in used:
                    f = next((x for x in FAMILIES if x not in used and x not in
                              {PREFERRED.get(y) for y in shapes if y != s}), "grey")
                fam[s] = f; used.add(f)
            ordered = [v for s in shapes for v in sorted((x for x in ms if shape_of(x) == s), key=lambda x: x["short_name"])]
            head = ["Group (slots × context)", "Configuration", "Average", "Slowest level", "Fastest level", "Up", "Down"]
            data.cell(row=row, column=1, value=f"{title}: {mname}").font = Font(name=FONT, bold=True)
            for i, h in enumerate(head, 1):
                c = data.cell(row=row + 1, column=i, value=h)
                c.font = Font(name=FONT, bold=True, color="FFFFFF"); c.fill = HEAD_FILL
            r0 = row + 2
            for i, v in enumerate(ordered):
                r = r0 + i
                crit = f'{rng("short_name")},$B{r},{rng(m)},"<>"'
                data.cell(row=r, column=1, value=shape_of(v)); data.cell(row=r, column=2, value=v["short_name"])
                data.cell(row=r, column=3, value=f"=AVERAGEIFS({rng(m)},{crit})")
                data.cell(row=r, column=4, value=f"=_xlfn.MINIFS({rng(m)},{crit})")
                data.cell(row=r, column=5, value=f"=_xlfn.MAXIFS({rng(m)},{crit})")
                data.cell(row=r, column=6, value=f"=E{r}-C{r}")
                data.cell(row=r, column=7, value=f"=C{r}-D{r}")
                for cc in range(3, 8):
                    data.cell(row=r, column=cc).number_format = "#,##0.0"
            r1 = r0 + len(ordered) - 1
            ch = BarChart()
            ch.type = "bar"                           # horizontal bars: the long names fit
            ch.title = f"{title}: {mname}"
            ch.style = 10
            ch.legend = None
            ch.height = max(7, 0.62 * len(ordered) + 2.5)
            ch.width = 26
            ch.gapWidth = 40
            ch.y_axis.title = mname
            ch.y_axis.majorGridlines = None
            ch.y_axis.scaling.min = 0                 # bars from zero: lengths compare honestly
            ch.y_axis.delete = False
            ch.x_axis.delete = False
            ch.x_axis.scaling.orientation = "maxMin"   # first group on top
            ref = Reference(data, min_col=3, min_row=r0, max_row=r1)
            ch.add_data(ref, titles_from_data=False)
            ch.set_categories(Reference(data, min_col=2, min_row=r0, max_row=r1))
            se = ch.series[0]
            shade = {}
            for i, v in enumerate(ordered):
                s = shape_of(v)
                k = shade.get(s, 0); shade[s] = k + 1
                colr = FAMILIES[fam[s]][min(k, 5)]
                dp = DataPoint(idx=i)
                dp.graphicalProperties.solidFill = colr
                dp.graphicalProperties.line.solidFill = colr
                se.dPt.append(dp)
            se.errBars = ErrorBars(errBarType="both", errValType="cust", noEndCap=False,
                                   plus=NumDataSource(numRef=NumRef(f=f"'Chart data'!$F${r0}:$F${r1}")),
                                   minus=NumDataSource(numRef=NumRef(f=f"'Chart data'!$G${r0}:$G${r1}")),
                                   spPr=GraphicalProperties(ln=LineProperties(solidFill="000000", w=12700)))
            cs.add_chart(ch, f"A{chart_row}")
            chart_row += int(ch.height * 2) + 2       # ~2 rows per cm
            row = r1 + 3
    data.column_dimensions["A"].width = 26
    data.column_dimensions["B"].width = 46
    for c in "CDEFG":
        data.column_dimensions[c].width = 14


def readme(wb, res, counts, used_keys):
    ws = wb.active
    ws.title = "Read me"
    ws["A1"] = f"Benchmark results — {res.name}"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = (f"{counts['levels']} load levels, {counts['variants']} configurations, {counts['items']} queue items. "
                "Every table can be filtered, sorted and pivoted (click into it, Insert → PivotTable).")
    ws["A2"].font = Font(name=FONT, italic=True, color="555555")
    r = 4
    ws.cell(row=r, column=1, value="Sheets").font = Font(name=FONT, bold=True, size=12)
    for name, what in SHEETS:
        r += 1
        ws.cell(row=r, column=1, value=name).font = Font(name=FONT, bold=True)
        ws.cell(row=r, column=2, value=what).font = Font(name=FONT)
    r += 2
    ws.cell(row=r, column=1, value="Terms").font = Font(name=FONT, bold=True, size=12)
    for term, what in GLOSSARY:
        r += 1
        ws.cell(row=r, column=1, value=term).font = Font(name=FONT, bold=True)
        ws.cell(row=r, column=2, value=what).font = Font(name=FONT)
    r += 2
    ws.cell(row=r, column=1, value="Column names").font = Font(name=FONT, bold=True, size=12)
    ws.cell(row=r, column=2, value="Readable header → the field name in the files (CSV/JSON) and plan files").font = \
        Font(name=FONT, italic=True, color="555555")
    for k in sorted(used_keys, key=lambda k: label(k).lower()):
        if label(k) != k:
            r += 1
            ws.cell(row=r, column=1, value=label(k)).font = Font(name=FONT)
            ws.cell(row=r, column=2, value=k).font = Font(name="Courier New")
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 120
    for row in ws.iter_rows(min_row=5):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("results")
    ap.add_argument("-o", "--out")
    a = ap.parse_args(argv)
    res = Path(a.results)
    pv = plan_variants()
    att = lambda vid: attrs(vid, (pv.get(vid) or (None, None))[1])   # noqa: E731
    levels, items, tasks, box_rows, starts, variants = [], [], [], [], {}, {}
    raw_tp, raw_res, raw_ev = [], [], []
    box_long = read_box(res)
    for run in sorted(p for p in res.iterdir() if p.is_dir()):
        plan = run.name.removesuffix("-smoke")
        st = json.loads((run / "state.json").read_text())["items"] if (run / "state.json").exists() else []
        for it in st:
            vid = variant_of(it["id"])
            fill = next((x.split("=")[1] for x in it["id"].split("/") if x.startswith("fill=")), None)
            items.append({"run": run.name, "variant": vid, "short_name": att(vid)["short_name"], "item": it["id"],
                          "test": test_name(it["workload"], num(fill)), "status": it["status"],
                          "duration_s": it.get("duration_s"), "attempts": it.get("attempts"),
                          "started": it.get("started"), "finished": it.get("finished"),
                          "error": it.get("last_error") or None})
            variants.setdefault((run.name, vid), {"run": run.name, "variant": vid, **att(vid)})
        if (run / "throughput.csv").exists():
            for r in csv.DictReader(open(run / "throughput.csv")):
                raw_tp.append({"run": run.name, **{k: num(v) for k, v in r.items()}})
                vid = variant_of(r["label"])
                parts = r["label"].split("/")
                fill = num(next((x.split("=")[1] for x in parts if x.startswith("fill=")), None))
                row = {"run": run.name, "plan": plan, "variant": vid, "workload": parts[1] if len(parts) > 1 else "",
                       "fill": fill, "clients": num(r["streams"]), **att(vid)}
                row["test"] = test_name(row["workload"], fill)
                row.update({k: num(r.get(k)) for k in METRIC_COLS if k in r and k not in row})
                row["first_error"] = r.get("first_error") or None
                levels.append(row)
        if (run / "results.jsonl").exists():
            for l in (run / "results.jsonl").read_text().splitlines():
                if l.strip():
                    j = json.loads(l)
                    tasks.append({"run": run.name, "variant": variant_of(j["item"]),
                                  "short_name": att(variant_of(j["item"]))["short_name"], **j})
                    raw_res.append({"run": run.name, **{k: (json.dumps(v) if isinstance(v, (dict, list)) else v)
                                                        for k, v in j.items()}})
        if (run / "events.log").exists():
            for l in (run / "events.log").read_text().splitlines():
                ts, _, msg = l.partition(" ")
                raw_ev.append({"run": run.name, "time": ts, "event": msg})
        known = {variant_of(it["id"]) for it in st}
        for f in run.glob("up-*.log"):
            vid = f.stem[3:]
            if known and vid not in known:              # a stale file from an earlier run in the same folder
                continue
            starts[(run.name, vid)] = start_info(f)
        if not box_long and (run / "report.md").exists():
            for b in box_from_report(run / "report.md"):
                box_rows.append({"run": run.name, "variant": variant_of(str(b.get("item", ""))), **b})
    if box_long:                           # average the raw samples per item window
        keys = sorted({k for _, k, _ in box_long})
        for it in items:
            a0, b0 = it.get("started"), it.get("finished")
            if not a0 or not b0:
                continue
            row = {"run": it["run"], "short_name": it["short_name"], "item": it["item"], "test": it["test"],
                   "s": round(b0 - a0)}
            for k in keys:
                v = [x for t_, kk, x in box_long if kk == k and a0 <= t_ <= b0 + 5]
                if v:
                    row[k] = round(max(v) if k.endswith(("mem_mb", "used_gb")) else statistics.mean(v), 1)
            box_rows.append(row)
    for key, s in starts.items():
        variants.setdefault(key, {"run": key[0], "variant": key[1], **att(key[1])})
        variants[key].update(s)
    vrows = sorted(variants.values(), key=lambda r: (r["role"], r["engine"], r["short_name"]))
    dup = {s for s in (v["short_name"] for v in vrows) if sum(1 for w in vrows if w["short_name"] == s) > 1}
    if dup:                                        # short names must be unique (the Summary's formulas key on them)
        for coll in [vrows, levels, items, tasks, box_rows]:
            for r in coll:
                if r.get("short_name") in dup:
                    r["short_name"] += f" [{r['variant'] if 'variant' in r else variant_of(r['item'])}]"
    levels.sort(key=lambda r: (r["role"], r["short_name"], r["test"], r["clients"] or 0))

    wb = Workbook()
    used = set(LEVEL_COLS) | {k for r in raw_tp + raw_res for k in r} | {k for _, k, _ in box_long}
    readme(wb, res, {"levels": len(levels), "variants": len(vrows), "items": len(items)}, used)
    sm = wb.create_sheet("Summary")
    data_sheet(wb, "Load levels", LEVEL_COLS, levels, widths={"short_name": 34, "variant": 30, "test": 30,
                                                                 "configuration": 90, "first_error": 40})
    var_cols = (["short_name", "variant", "configuration"] + ATTR_COLS + ["fallback", "started", "ready_s",
                "fail_reason", "strata_experts_in_vram", "strata_expert_vram_gib", "strata_experts_ram_gib",
                "strata_expert_load_gibs", "vllm_kv_tokens", "vllm_kv_gib", "vllm_max_concurrency", "llama_slots",
                "llama_ctx_slot", "run"])
    data_sheet(wb, "Configurations", var_cols, vrows, widths={"short_name": 34, "variant": 30, "configuration": 90,
                                                             "fail_reason": 60})
    data_sheet(wb, "Queue items", ["short_name", "test", "status", "duration_s", "attempts", "error", "item", "run",
                                   "started", "finished"], items, widths={"short_name": 34, "test": 30, "item": 60,
                                                                          "error": 50})
    tcols = ["short_name", "task", "shape", "level", "rep", "ok", "passed", "finish", "prompt_tokens",
             "completion_tokens", "reasoning_tokens", "reasoning_src", "cached_tokens", "ttft_s", "think_s", "total_s",
             "item", "run"]
    data_sheet(wb, "Answer tasks", tcols, tasks, widths={"short_name": 34, "item": 60})
    bkeys = sorted({k for r in box_rows for k in r} - {"run", "short_name", "item", "test", "s", "variant"},
                   key=lambda k: (not k.startswith("gpu"), k))
    data_sheet(wb, "Box per item", ["short_name", "test", "s"] + bkeys + ["item", "run"], box_rows,
               widths={"short_name": 34, "test": 30, "item": 60})

    def cols_of(rows, first):
        seen = list(first)
        for r in rows:
            seen += [k for k in r if k not in seen]
        return seen
    if raw_tp:
        data_sheet(wb, "Raw load levels", cols_of(raw_tp, ["run", "label", "streams"]), raw_tp, widths={"label": 55})
    if raw_res:
        data_sheet(wb, "Raw requests", cols_of(raw_res, ["run", "item"]), raw_res, widths={"item": 55})
    if box_long:
        import datetime as _dt
        bk = sorted({k for _, k, _ in box_long}, key=lambda k: (not k.startswith("gpu"), k))
        by_t = {}
        for ts, k, v in box_long:
            by_t.setdefault(ts, {})[k] = v
        rows_b = [{"epoch": ts, "time": _dt.datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M:%S"), **by_t[ts]}
                  for ts in sorted(by_t)]
        data_sheet(wb, "Raw box samples", ["time", "epoch"] + bk, rows_b, widths={"time": 19})
    if raw_ev:
        data_sheet(wb, "Raw events", ["run", "time", "event"], raw_ev, widths={"event": 110})

    # ---------------------------------------------------------------- Summary: formulas over "Load levels" + charts
    sm["A1"] = "Summary — one row per configuration"
    sm["A1"].font = Font(name=FONT, bold=True, size=14)
    sm["A2"] = ("Numbers are formulas over \"Load levels\" (averages, best or total over the configuration's levels). "
                "Empty = not measured for this configuration (or it did not start: see \"Configurations\"). "
                "Terms: \"Read me\".")
    sm["A2"].font = Font(name=FONT, italic=True, color="555555")
    hdr = ["Short name", "Configuration", "Start time (s)"] + [h for h, *_ in SUMMARY]
    for i, h in enumerate(hdr, 1):
        c = sm.cell(row=4, column=i, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF"); c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    sm.row_dimensions[4].height = 75
    n = len(levels) + 1
    col = {c: get_column_letter(LEVEL_COLS.index(c) + 1) for c in LEVEL_COLS}
    rng = lambda c: f"'Load levels'!${col[c]}$2:${col[c]}${n}"      # noqa: E731
    vstart = 5
    fn = {"avg": "AVERAGEIFS", "max": "_xlfn.MAXIFS", "sum": "SUMIFS"}   # _xlfn.: Excel 2019+ function
    for i, v in enumerate(vrows):
        r = vstart + i
        sm.cell(row=r, column=1, value=v["short_name"]); sm.cell(row=r, column=2, value=v["configuration"])
        sm.cell(row=r, column=3, value=v.get("ready_s"))
        for k, (h, kind, m, test) in enumerate(SUMMARY):
            crit = f'{rng("short_name")},$A{r}' + (f',{rng("test")},"{test}"' if test else "")
            cnt = f'COUNTIFS({crit},{rng(m)},"<>")'
            f = f'=IF({cnt}=0,"",{fn[kind]}({rng(m)},{crit}))'
            c = sm.cell(row=r, column=4 + k, value=f)
            c.number_format = "0.00" if "share" in h else "#,##0.0"
        for cell in sm[r]:
            cell.font = Font(name=FONT)
    vend = vstart + len(vrows) - 1
    sm.column_dimensions["A"].width = 38
    sm.column_dimensions["B"].width = 60
    for i in range(3, len(hdr) + 1):
        sm.column_dimensions[get_column_letter(i)].width = 15
    for r in range(vstart, vend + 1):
        sm.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    charts = [("Prompt reading per request (tok/s)", [SUMMARY[0][0], SUMMARY[3][0]]),
              ("Writing speed per request, steady (tok/s)", [SUMMARY[1][0], SUMMARY[5][0]]),
              ("Writing speed, all requests at the best level (tok/s)", [SUMMARY[2][0]])]
    anchor_row = vend + 3
    for k, (title, series) in enumerate(charts):
        ch = BarChart()
        ch.type = "bar"
        ch.title = title
        ch.style = 10
        ch.height = max(8, 0.55 * len(vrows) + 3)
        ch.width = 22
        for s in series:
            ci = hdr.index(s) + 1
            ch.add_data(Reference(sm, min_col=ci, min_row=4, max_row=vend), titles_from_data=True)
        ch.set_categories(Reference(sm, min_col=1, min_row=vstart, max_row=vend))
        ch.legend.position = "b"
        sm.add_chart(ch, f"{get_column_letter(1 + 7 * k) if k else 'A'}{anchor_row}")
    sm.freeze_panes = "C5"

    charts_sheet(wb, vrows, levels, rng)

    out = Path(a.out or res / "bench-results.xlsx")
    wb.save(out)
    print(f"wrote {out} — {len(levels)} levels, {len(vrows)} configurations, {len(items)} items, {len(tasks)} task rows, "
          f"{len(box_rows)} box rows")
    return 0


if __name__ == "__main__":
    sys.exit(main())
